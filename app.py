from pathlib import Path
import math
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from openpyxl import load_workbook

# AKBAL | Yacimientos I
# Aplicación de consulta: las etapas y los datos originales son de solo lectura.

st.set_page_config(page_title="AKBAL | Yacimientos I", page_icon="🛢️", layout="wide")

BLUE = "#0B1F3A"
TEAL = "#19A7A1"
EXPECTED_STAGES = [0, 2, 5, 6, 7, 8]
BOOK_PATH = Path(__file__).parent / "Excel AKBAL.xlsx"

st.markdown(f"""
<style>
  .stApp {{ background: #F4F7FB; }}
  [data-testid="stSidebar"] {{ background: {BLUE}; }}
  [data-testid="stSidebar"] * {{ color: #F4F7FB !important; }}
  h1, h2, h3 {{ color: {BLUE}; }}
  div[data-testid="stMetric"] {{ background: white; border: 1px solid #D9E2EF; padding: 14px; border-radius: 12px; }}
  .akbal-note {{ background: #E8F1FA; padding: 12px 16px; border-left: 4px solid {TEAL}; border-radius: 6px; }}
</style>
""", unsafe_allow_html=True)

@st.cache_data

def load_source(path_string: str, modified: float):
    """Lee solamente celdas fuente; no usa la interfaz para editar el historial."""
    wb = load_workbook(path_string, data_only=True, read_only=True)
    required = {"01_PARAMETROS", "01_CARACTERIZACION", "02_PVT", "03_HISTORIAL", "04_EXPANSION", "05_RETIRO"}
    missing = required.difference(wb.sheetnames)
    if missing:
        raise ValueError(f"Faltan hojas requeridas en el libro: {', '.join(sorted(missing))}")

    p = wb["01_PARAMETROS"]
    params = {
        "Vb": float(p["B3"].value), "phi": float(p["B4"].value), "Swi": float(p["B5"].value),
        "Pi": float(p["B6"].value), "Pb": float(p["B7"].value), "T": float(p["B8"].value),
        "Boi": float(p["B9"].value), "Rsi": float(p["B10"].value),
        "Caq": float(p["B11"].value), "Bw": float(p["B12"].value),
    }
    pvt = wb["02_PVT"]
    hist = wb["03_HISTORIAL"]
    rows = []
    for r in range(4, 10):
        stage = int(pvt.cell(r, 1).value)
        # Las celdas de datos fuente del historial son literales en el libro.
        pressure = float(pvt.cell(r, 2).value)
        bo = float(pvt.cell(r, 3).value)
        rs = float(pvt.cell(r, 4).value)
        bg_raw = pvt.cell(r, 5).value
        bg = None if bg_raw is None or (isinstance(bg_raw, str) and not bg_raw.strip()) else float(bg_raw)
        np_ = float(hist.cell(r, 3).value)
        gp = float(hist.cell(r, 4).value)
        wp = float(hist.cell(r, 5).value)
        rows.append({"Etapa": stage, "Presión (kg/cm²(a))": pressure, "Bo (m³/m³)": bo,
                     "Rs (m³/m³)": rs, "Bg (m³/m³)": bg, "Np (m³)": np_, "Gp (m³)": gp, "Wp (m³)": wp})
    wb.close()
    stages = [x["Etapa"] for x in rows]
    if stages != EXPECTED_STAGES:
        raise ValueError(f"Las etapas del archivo no coinciden con el inventario aprobado. Encontradas: {stages}; esperadas: {EXPECTED_STAGES}. No se reordenarán ni corregirán automáticamente.")
    if not math.isclose(params["Pi"], rows[0]["Presión (kg/cm²(a))"], abs_tol=1e-9):
        raise ValueError("La presión inicial Pi no coincide con la presión de la primera etapa. Revisa el libro fuente antes de continuar.")
    return params, pd.DataFrame(rows)

st.sidebar.markdown("# AKBAL")
st.sidebar.caption("Yacimientos I · Modelo de consulta")
page = st.sidebar.radio("Navegación", ["Panel principal", "Caracterización", "Análisis PVT", "Historial de producción", "Expansión de fluidos", "Entrada de agua", "Retiro subterráneo", "Balance de materia", "Dashboard", "Modelo 3D del yacimiento"])
st.sidebar.markdown("---")
st.sidebar.caption("Etapas bloqueadas: 0 · 2 · 5 · 6 · 7 · 8")

try:
    params, df = load_source(str(BOOK_PATH), BOOK_PATH.stat().st_mtime)
except Exception as e:
    st.error(f"No se pudo validar el libro fuente: {e}")
    st.stop()

# Cálculos reproducidos en Python a partir de las ecuaciones del libro.
PV = params["Vb"] * params["phi"]
Vw = PV * params["Swi"]
HCPV = PV * (1 - params["Swi"])
closure = Vw + HCPV - PV

df["Estado"] = df["Presión (kg/cm²(a))"].apply(lambda p: "Subsaturado" if p > params["Pb"] else ("Punto de burbuja" if math.isclose(p, params["Pb"], abs_tol=1e-9) else "Bifásico"))
df["Rp (m³/m³)"] = df.apply(lambda r: (r["Gp (m³)"] / r["Np (m³)"]) if r["Np (m³)"] else None, axis=1)
df["ΔP (kg/cm²)"] = params["Pi"] - df["Presión (kg/cm²(a))"]
df["Agotamiento (%)"] = df["ΔP (kg/cm²)"] / params["Pi"] * 100
# Eo = Bo-Boi por encima/en Pb; por debajo, añade (Rsi-Rs)*Bg.
def calc_eo(r):
    base = r["Bo (m³/m³)"] - params["Boi"]
    if r["Presión (kg/cm²(a))"] >= params["Pb"]:
        return base
    if r["Bg (m³/m³)"] is None:
        return None
    return base + (params["Rsi"] - r["Rs (m³/m³)"]) * r["Bg (m³/m³)"]
df["Eo (m³/m³)"] = df.apply(calc_eo, axis=1)
df["We (m³)"] = params["Caq"] * df["ΔP (kg/cm²)"]
df["We·Bw (m³)"] = df["We (m³)"] * params["Bw"]
df["Np·Bo (m³)"] = df["Np (m³)"] * df["Bo (m³/m³)"]
df["Wp·Bw (m³)"] = df["Wp (m³)"] * params["Bw"]
def calc_f(r):
    base = r["Np·Bo (m³)"] + r["Wp·Bw (m³)"]
    if r["Presión (kg/cm²(a))"] >= params["Pb"]:
        return base
    bg = r["Bg (m³/m³)"]
    rp = r["Rp (m³/m³)"]
    if bg is None or rp is None:
        return None
    return base + r["Np (m³)"] * (rp - r["Rs (m³/m³)"]) * bg
df["F (m³)"] = df.apply(calc_f, axis=1)

# Branding
st.markdown("<div class='akbal-note'><b>AKBAL · Sistema de análisis de yacimientos</b><br>Datos históricos de solo lectura. Las etapas, su identificador y su orden proceden del Excel fuente.</div>", unsafe_allow_html=True)
st.title(page)

if page == "Panel principal":
    st.write("Resumen del modelo, parámetros iniciales y estado de la auditoría.")
    c1,c2,c3,c4 = st.columns(4)
    c1.metric("Etapas originales", len(df), "IDs: 0, 2, 5, 6, 7, 8")
    c2.metric("Presión inicial", f"{params['Pi']:,.0f} kg/cm²(a)")
    c3.metric("Presión de burbuja", f"{params['Pb']:,.0f} kg/cm²(a)")
    c4.metric("HCPV", f"{HCPV:,.0f} m³")
    st.subheader("Parámetros iniciales del modelo")
    pview = pd.DataFrame([
        ["Vb",params['Vb'],"m³","Volumen bruto"],["φ",params['phi'],"fracción","Porosidad"],
        ["Swi",params['Swi'],"fracción","Saturación inicial de agua"],["Pi",params['Pi'],"kg/cm²(a)","Presión inicial"],
        ["Pb",params['Pb'],"kg/cm²(a)","Presión de burbuja"],["T",params['T'],"°C","Temperatura"],
        ["Boi",params['Boi'],"m³/m³","Factor volumétrico inicial de aceite"],["Rsi",params['Rsi'],"m³/m³","Relación inicial gas-aceite disuelto"],
        ["Caq",params['Caq'],"m³/(kg/cm²)","Coeficiente lineal del acuífero"],["Bw",params['Bw'],"m³/m³","Factor volumétrico del agua"]],
        columns=["Parámetro","Valor","Unidad","Descripción"])
    st.dataframe(pview, hide_index=True, use_container_width=True)
    st.subheader("Estado de validación")
    st.success("Integridad de etapas: correcta; los IDs coinciden con el inventario aprobado.")
    if abs(closure) < 1e-9:
        st.success(f"Cierre volumétrico de caracterización correcto: error = {closure:.6g} m³.")
    st.info("N y RF no se calculan todavía: el libro no define una ecuación explícita de balance de materia para estimarlos.")

elif page == "Caracterización":
    st.write("Cálculos derivados directamente de los parámetros originales.")
    a,b,c = st.columns(3); a.metric("Volumen poroso · PV",f"{PV:,.2f} m³"); b.metric("Volumen de agua · Vw",f"{Vw:,.2f} m³"); c.metric("Volumen de hidrocarburos · HCPV",f"{HCPV:,.2f} m³")
    out = pd.DataFrame({"Variable":["PV","Vw","HCPV","Vw + HCPV","Error de cierre"],"Ecuación":["Vb × φ","PV × Swi","PV × (1 − Swi)","Vw + HCPV","(Vw + HCPV) − PV"],"Resultado":[PV,Vw,HCPV,Vw+HCPV,closure],"Unidad":["m³"]*5})
    st.dataframe(out,hide_index=True,use_container_width=True)

elif page == "Análisis PVT":
    st.write("Tabla PVT de solo lectura. Bg = 0 en las etapas iniciales se conserva como dato del Excel; no se interpreta como un valor físico de gas calculado.")
    view = df[["Etapa","Presión (kg/cm²(a))","Bo (m³/m³)","Rs (m³/m³)","Bg (m³/m³)","Estado"]].copy()
    view["Bg (m³/m³)"] = view.apply(lambda r: "No aplicable / 0 en Excel" if r["Presión (kg/cm²(a))"] >= params["Pb"] and r["Bg (m³/m³)"] == 0 else r["Bg (m³/m³)"],axis=1)
    st.dataframe(view,hide_index=True,use_container_width=True)
    plot = df.melt(id_vars=["Etapa","Presión (kg/cm²(a))"],value_vars=["Bo (m³/m³)","Rs (m³/m³)","Bg (m³/m³)"],var_name="Propiedad",value_name="Valor")
    prop = st.selectbox("Propiedad para graficar",["Bo (m³/m³)","Rs (m³/m³)","Bg (m³/m³)"])
    fig=px.line(df,x="Presión (kg/cm²(a))",y=prop,markers=True,text="Etapa",title=f"{prop} vs presión")
    fig.update_xaxes(autorange="reversed"); fig.update_layout(template="plotly_white",paper_bgcolor="rgba(0,0,0,0)")
    st.plotly_chart(fig,use_container_width=True)
    st.caption("Rs y Bg pueden requerir una interpretación por región de presión. Los valores cero del archivo no se sustituyen ni se interpolan.")

elif page == "Historial de producción":
    st.write("Historial inmutable leído del Excel. Rp se reproduce como Gp/Np cuando Np es distinto de cero.")
    st.dataframe(df[["Etapa","Presión (kg/cm²(a))","Np (m³)","Gp (m³)","Wp (m³)","Rp (m³/m³)"]],hide_index=True,use_container_width=True)
    fig=px.line(df,x="Etapa",y=["Np (m³)","Gp (m³)","Wp (m³)"],markers=True,title="Producción acumulada por etapa")
    fig.update_layout(template="plotly_white",paper_bgcolor="rgba(0,0,0,0)"); st.plotly_chart(fig,use_container_width=True)
    chosen=st.selectbox("Consultar etapa original (solo lectura)",EXPECTED_STAGES,index=len(EXPECTED_STAGES)-1)
    row=df[df["Etapa"]==chosen].iloc[0]
    st.subheader(f"Detalle de la etapa {chosen}")
    cols=st.columns(3)
    for col,label,key in zip(cols,["Np","Gp","Wp"],["Np (m³)","Gp (m³)","Wp (m³)"]): col.metric(label,f"{row[key]:,.2f} m³")

elif page == "Expansión de fluidos":
    st.write("Eo se calcula con la ecuación del libro: por encima/en Pb, Bo − Boi; por debajo de Pb, se añade (Rsi − Rs) × Bg.")
    st.dataframe(df[["Etapa","Presión (kg/cm²(a))","Estado","Bo (m³/m³)","Rs (m³/m³)","Bg (m³/m³)","Eo (m³/m³)"]],hide_index=True,use_container_width=True)
    fig=px.line(df,x="Presión (kg/cm²(a))",y="Eo (m³/m³)",markers=True,text="Etapa",title="Expansión del sistema vs presión")
    fig.update_xaxes(autorange="reversed"); fig.update_layout(template="plotly_white",paper_bgcolor="rgba(0,0,0,0)"); st.plotly_chart(fig,use_container_width=True)

elif page == "Entrada de agua":
    st.write("Modelo lineal del Excel: We = Caq × (Pi − P). El equivalente volumétrico se calcula como We × Bw.")
    st.dataframe(df[["Etapa","Presión (kg/cm²(a))","ΔP (kg/cm²)","We (m³)","We·Bw (m³)"]],hide_index=True,use_container_width=True)
    fig=px.line(df,x="Etapa",y=["We (m³)","We·Bw (m³)"],markers=True,title="Entrada de agua por etapa")
    fig.update_layout(template="plotly_white",paper_bgcolor="rgba(0,0,0,0)"); st.plotly_chart(fig,use_container_width=True)
    st.caption("El acuífero lineal es una hipótesis del modelo didáctico, no una caracterización dinámica completa del acuífero.")

elif page == "Retiro subterráneo":
    st.write("Se reproduce el retiro F del Excel. Para P ≥ Pb: F = Np·Bo + Wp·Bw. Para P < Pb se añade Np·(Rp − Rs)·Bg.")
    st.dataframe(df[["Etapa","Estado","Np·Bo (m³)","Wp·Bw (m³)","F (m³)"]],hide_index=True,use_container_width=True)
    fig=px.line(df,x="Etapa",y="F (m³)",markers=True,title="Retiro subterráneo F por etapa")
    fig.update_layout(template="plotly_white",paper_bgcolor="rgba(0,0,0,0)"); st.plotly_chart(fig,use_container_width=True)

elif page == "Balance de materia":
    st.warning("No se calcularán N ni RF hasta que se establezca y documente la ecuación de balance de materia que corresponde al proyecto.")
    st.markdown("**Datos disponibles:** historial de Np, Gp, Wp; PVT por etapa; Eo; We; F; parámetros iniciales del yacimiento.")
    st.markdown("**Pendiente técnico:** definir el modelo (por ejemplo, balance volumétrico de aceite con o sin influjo de agua), confirmar los términos y unidades, y especificar qué incógnita se resolverá por etapa.")
    st.dataframe(df[["Etapa","Presión (kg/cm²(a))","Eo (m³/m³)","We (m³)","F (m³)"]],hide_index=True,use_container_width=True)

elif page == "Dashboard":
    st.write("Comparación general de variables derivadas en las etapas originales.")
    c1,c2,c3,c4=st.columns(4)
    c1.metric("Np final",f"{df.iloc[-1]['Np (m³)']:,.0f} m³")
    c2.metric("Gp final",f"{df.iloc[-1]['Gp (m³)']:,.0f} m³")
    c3.metric("We final",f"{df.iloc[-1]['We (m³)']:,.0f} m³")
    c4.metric("F final",f"{df.iloc[-1]['F (m³)']:,.2f} m³")
    chart=st.selectbox("Variable para comparar",["Np (m³)","Gp (m³)","Wp (m³)","Eo (m³/m³)","We (m³)","F (m³)","Agotamiento (%)"])
    fig=px.bar(df,x="Etapa",y=chart,color="Estado",title=f"{chart} por etapa original",text_auto=".3s")
    fig.update_layout(template="plotly_white",paper_bgcolor="rgba(0,0,0,0)"); st.plotly_chart(fig,use_container_width=True)
    st.subheader("Conclusiones técnicas preliminares")
    st.markdown(f"- La presión disminuye de {params['Pi']:,.0f} a {df.iloc[-1]['Presión (kg/cm²(a))']:,.0f} kg/cm²(a) a lo largo de las etapas registradas.\n- La presión de burbuja del modelo es {params['Pb']:,.0f} kg/cm²(a); las etapas 6, 7 y 8 están en región bifásica.\n- El influjo de agua calculado con el modelo lineal aumenta hasta {df.iloc[-1]['We (m³)']:,.0f} m³.\n- La estimación de N y RF queda pendiente de definir y validar la ecuación de balance de materia.")


elif page == "Modelo 3D del yacimiento":
    st.write("Exploración tridimensional esquemática del volumen del yacimiento y sus zonas conceptuales.")
    st.info(
        "Representación conceptual, no un modelo geológico real. El libro fuente aporta volumen bruto, "
        "porosidad y saturación inicial, pero no coordenadas, dimensiones geométricas, espesor, profundidad "
        "ni contacto aceite-agua medido. Por ello, la forma y la separación visual de las zonas son ilustrativas "
        "y no están a escala; no deben interpretarse como contactos o límites geológicos reales."
    )

    stage_3d = st.selectbox(
        "Consultar etapa original (solo lectura)",
        EXPECTED_STAGES,
        index=len(EXPECTED_STAGES) - 1,
        key="modelo3d_etapa_original",
    )
    row3d = df.loc[df["Etapa"] == stage_3d].iloc[0]

    # Coordenadas normalizadas sin unidades: geometría ilustrativa, no dimensiones del yacimiento.
    # Se usan prismas conceptuales para visualizar límites y zonas; no se infieren contactos reales.
    def add_prism(fig, x0, x1, y0, y1, z0, z1, color, name, opacity=0.38, showlegend=True):
        vertices = [
            (x0, y0, z0), (x1, y0, z0), (x1, y1, z0), (x0, y1, z0),
            (x0, y0, z1), (x1, y0, z1), (x1, y1, z1), (x0, y1, z1),
        ]
        xs, ys, zs = zip(*vertices)
        # Triangulación de las seis caras del prisma.
        i = [0, 0, 4, 4, 0, 0, 1, 1, 2, 2, 3, 3]
        j = [1, 2, 5, 6, 1, 5, 2, 6, 3, 7, 0, 4]
        k = [2, 3, 6, 7, 5, 4, 6, 5, 7, 6, 4, 7]
        fig.add_trace(go.Mesh3d(
            x=xs, y=ys, z=zs, i=i, j=j, k=k,
            color=color, opacity=opacity, name=name, legendgroup=name,
            showlegend=showlegend, hoverinfo="skip", flatshading=True,
        ))

    fig3d = go.Figure()
    # La zona de aceite y la zona de agua se separan solo para comunicar el concepto.
    add_prism(fig3d, 0, 10, 0, 6, 3.6, 8, "#19A7A1", "Zona conceptual asociada a aceite", 0.42)
    add_prism(fig3d, 0, 10, 0, 6, 0, 2.4, "#2684FF", "Zona conceptual asociada a agua", 0.42)

    # Marco externo como límites esquemáticos del volumen representado.
    corners = {
        "A": (0, 0, 0), "B": (10, 0, 0), "C": (10, 6, 0), "D": (0, 6, 0),
        "E": (0, 0, 8), "F": (10, 0, 8), "G": (10, 6, 8), "H": (0, 6, 8),
    }
    edges = [("A","B"),("B","C"),("C","D"),("D","A"),("E","F"),("F","G"),("G","H"),("H","E"),("A","E"),("B","F"),("C","G"),("D","H")]
    for a, b in edges:
        fig3d.add_trace(go.Scatter3d(
            x=[corners[a][0], corners[b][0]], y=[corners[a][1], corners[b][1]], z=[corners[a][2], corners[b][2]],
            mode="lines", line=dict(color="#0B1F3A", width=5),
            name="Límite esquemático del volumen", legendgroup="Límite esquemático del volumen",
            showlegend=(a, b) == edges[0], hoverinfo="skip",
        ))

    fig3d.update_layout(
        template="plotly_white", height=560,
        margin=dict(l=0, r=0, t=45, b=0),
        title="Modelo 3D conceptual · límites y zonas ilustrativas",
        scene=dict(
            xaxis=dict(title="Eje X normalizado", showbackground=True, backgroundcolor="rgba(11,31,58,0.04)", gridcolor="#D9E2EF"),
            yaxis=dict(title="Eje Y normalizado", showbackground=True, backgroundcolor="rgba(11,31,58,0.04)", gridcolor="#D9E2EF"),
            zaxis=dict(title="Eje Z normalizado", showbackground=True, backgroundcolor="rgba(11,31,58,0.04)", gridcolor="#D9E2EF"),
            aspectmode="manual", aspectratio=dict(x=1.5, y=0.9, z=0.8),
            camera=dict(eye=dict(x=1.6, y=1.6, z=1.25)),
        ),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
    )
    st.plotly_chart(fig3d, use_container_width=True, config={"responsive": True, "displaylogo": False, "scrollZoom": True})
    st.caption("Colores: turquesa = zona conceptual asociada a aceite; azul = zona conceptual asociada a agua; líneas azul marino = límites esquemáticos. Usa el mouse o los gestos táctiles para girar y acercar/alejar.")

    st.subheader(f"Indicadores de la etapa {stage_3d}")
    m1, m2, m3 = st.columns(3)
    m1.metric("Presión", f"{row3d['Presión (kg/cm²(a))']:,.2f} kg/cm²(a)")
    m2.metric("Estado calculado", str(row3d["Estado"]))
    m3.metric("Agotamiento", f"{row3d['Agotamiento (%)']:,.2f}%")
    m4, m5, m6 = st.columns(3)
    m4.metric("Np", f"{row3d['Np (m³)']:,.2f} m³")
    m5.metric("Gp", f"{row3d['Gp (m³)']:,.2f} m³")
    m6.metric("Wp", f"{row3d['Wp (m³)']:,.2f} m³")
    m7, m8, m9 = st.columns(3)
    m7.metric("Eo", "N/D" if pd.isna(row3d["Eo (m³/m³)"]) else f"{row3d['Eo (m³/m³)']:,.4f} m³/m³")
    m8.metric("We", f"{row3d['We (m³)']:,.2f} m³")
    m9.metric("F", "N/D" if pd.isna(row3d["F (m³)"]) else f"{row3d['F (m³)']:,.2f} m³")
    st.caption("Los indicadores proceden del historial y de las ecuaciones existentes en AKBAL. La selección solo consulta los datos; no modifica las etapas ni los cálculos.")

st.markdown("---")
st.caption("AKBAL · Prototipo técnico. Los datos de entrada y las etapas son de solo lectura. Valide el modelo con el responsable académico antes de usarlo para decisiones operativas.")
